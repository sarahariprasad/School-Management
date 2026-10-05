# students/utils.py
import openpyxl
from openpyxl.utils import get_column_letter
from django.core.mail import send_mail
from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone

STUDENT_EXPORT_HEADERS = [
    "Student ID", "Name", "Age", "Group", "Gender",
    "Date of Birth", "Admission Date", "Status", "Inactive Date",
    "Mother Name", "Mother Phone", "Mother Email",
    "Father Name", "Father Phone", "Father Email",
    "Address", "Active Therapies",
]


def _student_row(student):
    """Build one export row for a Student. Shared by the bulk and single-student
    exports so the columns never drift apart between the two."""
    active_therapies = ", ".join(
        f"{e.therapy.name} ({e.staff.full_name if e.staff else 'unassigned'})"
        for e in student.therapy_enrollments.filter(is_active=True)
    )
    return [
        student.student_id,
        student.name,
        student.age if student.age is not None else "",
        student.group.get_name_display() if student.group else "",
        student.gender,
        student.date_of_birth.strftime("%Y-%m-%d") if student.date_of_birth else "",
        student.admission_date.strftime("%Y-%m-%d") if student.admission_date else "",
        "Active" if student.is_active else "Inactive",
        student.inactive_date.strftime("%Y-%m-%d") if student.inactive_date else "",
        student.mother_name,
        student.mother_phone or "",
        student.mother_email or "",
        student.father_name,
        student.father_phone or "",
        student.father_email or "",
        student.address,
        active_therapies,
    ]


def _autosize(ws, header_count, width=20):
    for i in range(1, header_count + 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def export_students_to_excel(students):
    """Build an .xlsx HttpResponse listing every student in the given queryset."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    ws.append(STUDENT_EXPORT_HEADERS)
    for student in students:
        ws.append(_student_row(student))

    _autosize(ws, len(STUDENT_EXPORT_HEADERS))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="students_export.xlsx"'
    wb.save(response)
    return response


def export_single_student_to_excel(student):
    """Build an .xlsx HttpResponse for a single student's profile - a "Profile" sheet
    with their details, plus a "Status History" sheet and a "Remarks" sheet so the
    whole record (join date onward) is in one downloadable file."""
    wb = openpyxl.Workbook()

    profile_ws = wb.active
    profile_ws.title = "Profile"
    for label, value in zip(STUDENT_EXPORT_HEADERS, _student_row(student)):
        profile_ws.append([label, value])
    profile_ws.column_dimensions["A"].width = 22
    profile_ws.column_dimensions["B"].width = 40

    status_ws = wb.create_sheet("Status History")
    status_ws.append(["Status", "Since", "Reason"])
    for h in student.status_history.all().order_by("changed_on"):
        status_ws.append([
            h.get_status_display(),
            h.changed_on.strftime("%Y-%m-%d") if h.changed_on else "",
            h.reason or "",
        ])
    _autosize(status_ws, 3, width=22)

    remarks_ws = wb.create_sheet("Remarks")
    remarks_ws.append(["Date", "Category", "Remark", "Created By", "Sent to Parents"])
    for remark in student.remarks.all().order_by("created_at"):
        remarks_ws.append([
            remark.created_at.strftime("%Y-%m-%d %H:%M"),
            remark.get_category_display(),
            remark.remark_text,
            str(remark.created_by) if remark.created_by else "",
            "Yes" if remark.sent_to_parents else "No",
        ])
    for i, width in enumerate([18, 18, 60, 20, 14], start=1):
        remarks_ws.column_dimensions[get_column_letter(i)].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{student.student_id}_profile.xlsx"'
    wb.save(response)
    return response


def export_remarks_to_excel(student):
    """Export a single student's full remark/progress history to .xlsx (remarks only,
    no profile/status sheets - use export_single_student_to_excel for the full picture)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Remarks"

    ws.append(["Date", "Category", "Remark", "Created By", "Sent to Parents"])
    for remark in student.remarks.all().order_by("created_at"):
        ws.append([
            remark.created_at.strftime("%Y-%m-%d %H:%M"),
            remark.get_category_display(),
            remark.remark_text,
            str(remark.created_by) if remark.created_by else "",
            "Yes" if remark.sent_to_parents else "No",
        ])

    for i, width in enumerate([18, 18, 60, 20, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{student.student_id}_remarks.xlsx"'
    wb.save(response)
    return response


def send_remark_email(remark):
    """Email a progress remark to whichever parent email addresses are on file.
    Returns True if the email was sent, False if there was nothing to send to
    or sending failed (caller should surface a message in either case)."""
    student = remark.student
    recipients = [e for e in [student.mother_email, student.father_email] if e]

    if not recipients:
        return False

    subject = f"Progress Update for {student.name} ({student.student_id})"
    message = (
        f"Dear Parent,\n\n"
        f"Here is a progress update for {student.name}:\n\n"
        f"Category: {remark.get_category_display()}\n"
        f"Date: {remark.created_at:%d-%b-%Y}\n\n"
        f"{remark.remark_text}\n\n"
        f"Regards,\nSchool Administration"
    )

    try:
        send_mail(
            subject,
            message,
            getattr(settings, "DEFAULT_FROM_EMAIL", None),
            recipients,
            fail_silently=False,
        )
    except Exception:
        return False

    remark.sent_to_parents = True
    remark.sent_at = timezone.now()
    remark.save(update_fields=["sent_to_parents", "sent_at"])
    return True