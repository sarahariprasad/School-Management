from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Value
from django.db.models.functions import Concat
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.contenttypes.models import ContentType
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone

from accounts.decorators import role_required
from core.permissions import branch_scope
from audit_log.services import log_audit, snapshot_instance, build_field_changes, log_formset_changes
from audit_log.models import AuditLog
from .forms import (
    DocumentFormSet, EducationFormSet, EducationRecordForm,
    ExperienceFormSet, ExperienceHistoryForm, PromotionFormSet,
    PromotionHistoryForm, SalaryFormSet, SalaryIncrementForm,
    StaffCreateForm, StaffDocumentForm, StaffExitForm,
    StaffProfileForm, StaffSelfEditForm,
)
from .models import EducationRecord, StaffDocument, StaffProfile


# ═══════════════════════════════════════════════════════════════
# LIST VIEW
# ═══════════════════════════════════════════════════════════════
@login_required
def staff_list(request):
    search = request.GET.get("search", "").strip()
    status_filter = request.GET.get("status", "").strip().lower()

    staff = branch_scope(
        request.user,
        StaffProfile.objects.select_related(
            "user", "user__branch", "updated_by", "created_by"
        ),
        "user__branch",
    ).annotate(
        full_name=Concat("user__first_name", Value(" "), "user__last_name")
    )

    # ── Status filter ──
    if status_filter == "active":
        staff = staff.filter(is_active=True)
    elif status_filter == "inactive":
        staff = staff.filter(is_active=False)

    if search:
        staff = staff.filter(
            Q(employee_id__icontains=search) |
            Q(full_name__icontains=search) |
            Q(user__first_name__icontains=search) |
            Q(user__last_name__icontains=search) |
            Q(user__email__icontains=search) |
            Q(phone__icontains=search) |
            Q(designation__icontains=search) |
            Q(department__icontains=search) |
            Q(city__icontains=search) |
            Q(state__icontains=search)
        )

    paginator = Paginator(staff.order_by("employee_id"), 20)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    return render(request, "staff/list.html", {
        "staff": page_obj,
        "search": search,
        "status_filter": status_filter,
        "status_choices": [("active", "Active"), ("inactive", "Inactive")],
        "total_count": paginator.count,
    })


# ═══════════════════════════════════════════════════════════════
# CREATE VIEW  —  FIXED: validate everything BEFORE touching DB
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
@transaction.atomic
def staff_create(request):
    if request.method == "POST":
        form = StaffCreateForm(request.POST, request.FILES, actor=request.user)

        if form.is_valid():
            # 1) Build the profile in memory ONLY — do NOT write to DB yet.
            #    (This also prevents the User from being created early.)
            profile = form.save(commit=False)

            # 2) Bind formsets to the in-memory profile so they can validate
            education_formset = EducationFormSet(
                request.POST, request.FILES, instance=profile, prefix="edu"
            )
            document_formset = DocumentFormSet(
                request.POST, request.FILES, instance=profile, prefix="doc"
            )
            salary_formset = SalaryFormSet(
                request.POST, request.FILES, instance=profile, prefix="sal"
            )

            if (education_formset.is_valid() and document_formset.is_valid()
                    and salary_formset.is_valid()):
                # 3) EVERYTHING is clean — now actually persist
                form._save_with_user(profile)
                education_formset.save()
                document_formset.save()
                salary_formset.save()

                # ── AUDIT: main profile ──
                log_audit(request.user, profile, "CREATE",
                          snapshot=snapshot_instance(profile))

                # ── AUDIT: inline records ──
                for fs in (education_formset, document_formset, salary_formset):
                    for frm in fs.forms:
                        if frm.instance.pk and not frm.cleaned_data.get("DELETE"):
                            log_audit(
                                request.user, frm.instance, "CREATE",
                                snapshot=snapshot_instance(frm.instance)
                            )

                messages.success(
                    request,
                    f"Staff member '{profile.full_name}' created successfully."
                )
                return redirect("staff_list")
            else:
                messages.error(request, "Please correct the errors below.")
        else:
            # Main form invalid — still bind formsets so template can render POST data
            education_formset = EducationFormSet(
                request.POST, request.FILES, prefix="edu"
            )
            document_formset = DocumentFormSet(
                request.POST, request.FILES, prefix="doc"
            )
            salary_formset = SalaryFormSet(
                request.POST, request.FILES, prefix="sal"
            )
            messages.error(request, "Please correct the errors below.")
    else:
        form = StaffCreateForm(actor=request.user)
        education_formset = EducationFormSet(prefix="edu")
        document_formset = DocumentFormSet(prefix="doc")
        salary_formset = SalaryFormSet(prefix="sal")

    return render(request, "staff/staff_create_tabs.html", {
        "profile_form": form,
        "education_formset": education_formset,
        "document_formset": document_formset,
        "salary_formset": salary_formset,
        "title": "Add Staff Member",
        "action": "Create",
    })


# ═══════════════════════════════════════════════════════════════
# EDIT VIEW
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
@transaction.atomic
def staff_edit(request, pk):
    profile = get_object_or_404(
        branch_scope(
            request.user,
            StaffProfile.objects.select_related("user"),
            "user__branch",
        ),
        pk=pk,
    )

    # Capture a pristine copy BEFORE the form touches the instance
    old_profile = StaffProfile.objects.get(pk=profile.pk) if request.method == "POST" else None

    # Fetch audit trail for the template
    profile_ct = ContentType.objects.get_for_model(StaffProfile)
    audit_logs = AuditLog.objects.filter(
        content_type=profile_ct, object_id=profile.pk
    ).select_related("user").order_by("-action_time")[:25]

    if request.method == "POST":
        form = StaffProfileForm(
            request.POST, request.FILES, instance=profile, user=request.user
        )
        education_formset = EducationFormSet(
            request.POST, request.FILES, instance=profile, prefix="edu"
        )
        document_formset = DocumentFormSet(
            request.POST, request.FILES, instance=profile, prefix="doc"
        )
        salary_formset = SalaryFormSet(
            request.POST, request.FILES, instance=profile, prefix="sal"
        )
        experience_formset = ExperienceFormSet(
            request.POST, request.FILES, instance=profile, prefix="exp"
        )
        promotion_formset = PromotionFormSet(
            request.POST, request.FILES, instance=profile, prefix="promo"
        )

        if (form.is_valid() and education_formset.is_valid()
                and document_formset.is_valid() and salary_formset.is_valid()
                and experience_formset.is_valid() and promotion_formset.is_valid()):
            profile = form.save()
            education_formset.save()
            document_formset.save()
            salary_formset.save()
            experience_formset.save()
            promotion_formset.save()

            # ── AUDIT: Profile changes ──
            if old_profile:
                changes = build_field_changes(old_profile, profile)
                if changes:
                    log_audit(request.user, profile, "UPDATE", field_changes=changes)

            # ── AUDIT: Inline formsets ──
            log_formset_changes(request.user, education_formset, profile, "staff")
            log_formset_changes(request.user, document_formset, profile, "staff")
            log_formset_changes(request.user, salary_formset, profile, "staff")
            log_formset_changes(request.user, experience_formset, profile, "staff")
            log_formset_changes(request.user, promotion_formset, profile, "staff")

            messages.success(request, f"Staff '{profile.full_name}' updated successfully.")
            return redirect("staff_list")
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        form = StaffProfileForm(instance=profile, user=request.user)
        education_formset = EducationFormSet(instance=profile, prefix="edu")
        document_formset = DocumentFormSet(instance=profile, prefix="doc")
        salary_formset = SalaryFormSet(instance=profile, prefix="sal")
        experience_formset = ExperienceFormSet(instance=profile, prefix="exp")
        promotion_formset = PromotionFormSet(instance=profile, prefix="promo")

    return render(request, "staff/profile_edit.html", {
        "profile_form": form,
        "education_formset": education_formset,
        "document_formset": document_formset,
        "salary_formset": salary_formset,
        "experience_formset": experience_formset,
        "promotion_formset": promotion_formset,
        "audit_logs": audit_logs,
        "title": f"Edit Staff: {profile.full_name}",
        "action": "Update",
        "profile": profile,
    })


# ═══════════════════════════════════════════════════════════════
# DEACTIVATE VIEW
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
@transaction.atomic
def staff_deactivate(request, pk):
    profile = get_object_or_404(
        branch_scope(
            request.user,
            StaffProfile.objects.select_related("user"),
            "user__branch",
        ),
        pk=pk,
    )
    if not profile.is_active:
        return redirect("staff_list")

    old_profile = StaffProfile.objects.get(pk=profile.pk)
    form = StaffExitForm(request.POST or None, instance=profile)
    if request.method == "POST" and form.is_valid():
        profile = form.save(commit=False)
        profile.is_active = False
        profile.save(update_fields=["leaving_date", "exit_reason", "is_active"])
        profile.user.is_active = False
        profile.user.save(update_fields=["is_active"])

        # ── AUDIT ──
        changes = build_field_changes(old_profile, profile)
        log_audit(request.user, profile, "UPDATE", field_changes=changes)

        messages.success(request, f"Staff '{profile.full_name}' has been deactivated.")
        return redirect("staff_list")

    return render(request, "staff/deactivate.html", {
        "form": form,
        "profile": profile,
    })


# ═══════════════════════════════════════════════════════════════
# DOCUMENTS VIEW
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def staff_documents(request, pk):
    profile = get_object_or_404(
        branch_scope(
            request.user,
            StaffProfile.objects.select_related("user"),
            "user__branch",
        ),
        pk=pk,
    )
    return render(request, "staff/documents.html", {
        "profile": profile,
        "education_form": EducationRecordForm(),
        "document_form": StaffDocumentForm(),
    })


@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def education_add(request, pk):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    form = EducationRecordForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        record = form.save(commit=False)
        record.staff = profile
        record.save()
        log_audit(request.user, record, "CREATE", snapshot=snapshot_instance(record))
        messages.success(request, "Education record added.")
    return redirect("staff_documents", pk=pk)


@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def document_add(request, pk):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    form = StaffDocumentForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        document = form.save(commit=False)
        document.staff = profile
        document.save()
        log_audit(request.user, document, "CREATE", snapshot=snapshot_instance(document))
        messages.success(request, "Document uploaded.")
    return redirect("staff_documents", pk=pk)


# ═══════════════════════════════════════════════════════════════
# FILE DOWNLOAD
# ═══════════════════════════════════════════════════════════════
from django.http import FileResponse


@login_required
def staff_file_download(request, pk, kind, document_pk=None):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    if kind == "address":
        file_field = profile.address_proof
    elif kind == "education":
        file_field = get_object_or_404(EducationRecord, pk=document_pk, staff=profile).certificate
    else:
        file_field = get_object_or_404(StaffDocument, pk=document_pk, staff=profile).file

    if not file_field:
        messages.error(request, "File not found.")
        return redirect("staff_documents", pk=pk)

    return FileResponse(
        file_field.open("rb"),
        as_attachment=True,
        filename=file_field.name.rsplit("/", 1)[-1],
    )


# ═══════════════════════════════════════════════════════════════
# SELF PROFILE
# ═══════════════════════════════════════════════════════════════
@login_required
def staff_profile_view(request):
    profile = get_object_or_404(StaffProfile, user=request.user)
    return render(request, "staff/profile.html", {
        "profile": profile,
        "experience_history": profile.experience_history.all(),
        "promotions": profile.promotions.all(),
        "increments": profile.increments.all(),
        "is_self": True,
    })


@login_required
def staff_profile_edit(request):
    profile = get_object_or_404(StaffProfile, user=request.user)
    old_profile = StaffProfile.objects.get(pk=profile.pk) if request.method == "POST" else None

    form = StaffSelfEditForm(request.POST or None, request.FILES or None, instance=profile)
    if request.method == "POST" and form.is_valid():
        profile = form.save(commit=False)
        profile.updated_by = request.user
        profile.save()

        if old_profile:
            changes = build_field_changes(old_profile, profile)
            if changes:
                log_audit(request.user, profile, "UPDATE", field_changes=changes)

        messages.success(request, "Your profile has been updated.")
        return redirect("staff_profile")
    return render(request, "staff/self_edit.html", {"form": form, "profile": profile})


# ═══════════════════════════════════════════════════════════════
# ADMIN PROFILE VIEW
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def staff_profile_admin_view(request, pk):
    profile = get_object_or_404(
        branch_scope(
            request.user,
            StaffProfile.objects.select_related("user"),
            "user__branch",
        ),
        pk=pk,
    )

    # Optional: pass audit logs to the read-only profile template
    profile_ct = ContentType.objects.get_for_model(StaffProfile)
    audit_logs = AuditLog.objects.filter(
        content_type=profile_ct, object_id=profile.pk
    ).select_related("user").order_by("-action_time")[:15]

    return render(request, "staff/profile.html", {
        "profile": profile,
        "experience_history": profile.experience_history.all(),
        "promotions": profile.promotions.all(),
        "increments": profile.increments.all(),
        "audit_logs": audit_logs,
        "is_self": False,
    })


# ═══════════════════════════════════════════════════════════════
# HISTORY ADD VIEWS
# ═══════════════════════════════════════════════════════════════
@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def experience_add(request, pk):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    form = ExperienceHistoryForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        exp = form.save(commit=False)
        exp.staff = profile
        exp.save()
        log_audit(request.user, exp, "CREATE", snapshot=snapshot_instance(exp))
        messages.success(request, "Experience record added.")
        return redirect("staff_profile_admin_view", pk=profile.pk)
    return render(request, "staff/experience_form.html", {"form": form, "profile": profile})


@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def promotion_add(request, pk):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    form = PromotionHistoryForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        promo = form.save(commit=False)
        promo.staff = profile
        promo.save()
        log_audit(request.user, promo, "CREATE", snapshot=snapshot_instance(promo))
        messages.success(request, "Promotion record added.")
        return redirect("staff_profile_admin_view", pk=profile.pk)
    return render(request, "staff/promotion_form.html", {"form": form, "profile": profile})


@role_required("SYSTEM_ADMIN", "BRANCH_ADMIN")
def salary_increment_add(request, pk):
    profile = get_object_or_404(
        branch_scope(request.user, StaffProfile.objects.all(), "user__branch"),
        pk=pk,
    )
    form = SalaryIncrementForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        year = form.cleaned_data["year"]
        if profile.increments.filter(year=year).exists():
            form.add_error("year", "A salary increment for this year already exists.")
        else:
            inc = form.save(commit=False)
            inc.staff = profile
            if not inc.new_salary:
                inc.new_salary = inc.base_salary + inc.increment_amount
            inc.save()
            log_audit(request.user, inc, "CREATE", snapshot=snapshot_instance(inc))
            messages.success(request, "Salary increment added.")
            return redirect("staff_profile_admin_view", pk=profile.pk)
    return render(request, "staff/salary_increment_form.html", {"form": form, "profile": profile})


# ═══════════════════════════════════════════════════════════════
# EXCEL EXPORT — ALL STAFF
# ═══════════════════════════════════════════════════════════════
@login_required
def staff_export_excel(request):
    """Export all accessible staff to Excel."""
    staff_qs = branch_scope(
        request.user,
        StaffProfile.objects.select_related("user", "user__branch"),
        "user__branch",
    ).order_by("employee_id")

    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0d6efd", end_color="0d6efd", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "#", "Employee ID", "Full Name", "Email", "Gender", "Date of Birth",
        "Department", "Designation", "Branch", "Phone", "Emergency Contact",
        "Emergency Contact Name", "Joining Date", "Leaving Date", "Status",
        "Address", "City", "State", "Postal Code",
    ]
    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for idx, person in enumerate(staff_qs, start=1):
        ws.append([
            idx,
            person.employee_id,
            person.full_name,
            person.user.email,
            person.get_gender_display() or "—",
            person.date_of_birth.strftime("%d-%m-%Y") if person.date_of_birth else "—",
            person.get_department_display(),
            person.current_designation,
            person.user.branch.name if person.user.branch else "—",
            person.phone or "—",
            person.emergency_contact or "—",
            person.emergency_contact_name or "—",
            person.joining_date.strftime("%d-%m-%Y"),
            person.leaving_date.strftime("%d-%m-%Y") if person.leaving_date else "—",
            "Active" if person.is_active else "Inactive",
            person.address or "—",
            person.city or "—",
            person.state or "—",
            person.postal_code or "—",
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    column_widths = [5, 14, 25, 28, 10, 14, 16, 20, 18, 16, 16, 22, 14, 14, 10, 35, 14, 14, 14]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    import datetime
    filename = f"staff_list_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ═══════════════════════════════════════════════════════════════
# EXCEL EXPORT — SINGLE STAFF
# ═══════════════════════════════════════════════════════════════
@login_required
def staff_export_single_excel(request, pk):
    """Export a single staff member's complete details to Excel."""
    profile = get_object_or_404(
        branch_scope(
            request.user,
            StaffProfile.objects.select_related("user", "user__branch"),
            "user__branch",
        ),
        pk=pk,
    )

    wb = Workbook()

    # ── Sheet 1: Profile ──
    ws1 = wb.active
    ws1.title = "Profile"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0d6efd", end_color="0d6efd", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    profile_data = [
        ("Employee ID", profile.employee_id),
        ("Full Name", profile.full_name),
        ("Email", profile.user.email),
        ("Gender", profile.get_gender_display() or "—"),
        ("Date of Birth", profile.date_of_birth.strftime("%d-%m-%Y") if profile.date_of_birth else "—"),
        ("Department", profile.get_department_display()),
        ("Designation", profile.current_designation),
        ("Branch", profile.user.branch.name if profile.user.branch else "—"),
        ("Phone", profile.phone or "—"),
        ("Emergency Contact", profile.emergency_contact or "—"),
        ("Emergency Contact Name", profile.emergency_contact_name or "—"),
        ("Joining Date", profile.joining_date.strftime("%d-%m-%Y")),
        ("Leaving Date", profile.leaving_date.strftime("%d-%m-%Y") if profile.leaving_date else "—"),
        ("Status", "Active" if profile.is_active else "Inactive"),
        ("Address", profile.address or "—"),
        ("City", profile.city or "—"),
        ("State", profile.state or "—"),
        ("Postal Code", profile.postal_code or "—"),
        ("Exit Reason", profile.exit_reason or "—"),
    ]

    ws1.append(["Field", "Value"])
    for col in range(1, 3):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for label, value in profile_data:
        ws1.append([label, value])

    for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws1.column_dimensions["A"].width = 25
    ws1.column_dimensions["B"].width = 45

    # ── Sheet 2: Education ──
    ws2 = wb.create_sheet("Education")
    edu_headers = ["Qualification", "Institution", "Passing Year"]
    ws2.append(edu_headers)
    for col in range(1, 4):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for edu in profile.education_records.all():
        ws2.append([edu.qualification, edu.institution, edu.passing_year])

    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.border = thin_border

    for i in range(1, 4):
        ws2.column_dimensions[get_column_letter(i)].width = 25

    # ── Sheet 3: Experience ──
    ws3 = wb.create_sheet("Experience")
    exp_headers = ["Organization", "Role", "Start Date", "End Date"]
    ws3.append(exp_headers)
    for col in range(1, 5):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for exp in profile.experience_history.all():
        ws3.append([
            exp.organization,
            exp.role,
            exp.start_date.strftime("%d-%m-%Y"),
            exp.end_date.strftime("%d-%m-%Y") if exp.end_date else "Present",
        ])

    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = thin_border

    for i in range(1, 5):
        ws3.column_dimensions[get_column_letter(i)].width = 25

    # ── Sheet 4: Promotions ──
    ws4 = wb.create_sheet("Promotions")
    promo_headers = ["Old Designation", "New Designation", "Date", "Remarks"]
    ws4.append(promo_headers)
    for col in range(1, 5):
        cell = ws4.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for promo in profile.promotions.all():
        ws4.append([
            promo.old_designation,
            promo.new_designation,
            promo.promotion_date.strftime("%d-%m-%Y"),
            promo.remarks or "—",
        ])

    for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = thin_border

    for i in range(1, 5):
        ws4.column_dimensions[get_column_letter(i)].width = 25

    # ── Sheet 5: Salary ──
    ws5 = wb.create_sheet("Salary History")
    sal_headers = ["Year", "Base Salary", "Increment", "New Salary"]
    ws5.append(sal_headers)
    for col in range(1, 5):
        cell = ws5.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for sal in profile.increments.all():
        ws5.append([sal.year, sal.base_salary, sal.increment_amount, sal.new_salary])

    for row in ws5.iter_rows(min_row=1, max_row=ws5.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = thin_border

    for i in range(1, 5):
        ws5.column_dimensions[get_column_letter(i)].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    import datetime
    filename = f"staff_{profile.employee_id}_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response