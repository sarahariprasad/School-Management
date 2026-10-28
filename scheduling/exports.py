"""Excel (.xlsx) exports. Requires:  pip install openpyxl"""
from collections import Counter, defaultdict

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")


def _sheet(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 3, 42)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if hasattr(c.value, "year") and not hasattr(c.value, "hour"):
                c.number_format = "DD-MMM-YYYY"


def _respond(wb, filename):
    resp = HttpResponse(content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(resp)
    return resp


def _session_rows(qs):
    for s in qs:
        yield [
            s.date, s.date.strftime("%a"), str(s.time_slot), s.therapy.name, s.staff.full_name,
            s.student.student_id, s.student.name,
            s.student.group.name if s.student.group else "",
            s.get_status_display(), s.absence_reason,
            "Recurring" if s.assignment_id else "One-off",
        ]


SESSION_HEADERS = ["Date", "Day", "Time", "Therapy", "Therapist", "Student ID", "Student",
                   "Group", "Status", "Reason", "Type"]


def sessions_workbook(qs, filename, completion=None):
    qs = qs.select_related("student__group", "therapy", "staff__user", "time_slot")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    sessions = list(qs)
    _sheet(ws, SESSION_HEADERS, _session_rows(sessions))

    summary = defaultdict(Counter)
    for s in sessions:
        summary[s.staff.full_name][s.get_status_display()] += 1
    statuses = ["Scheduled", "Completed", "Student Absent", "Therapist Absent", "Cancelled"]
    _sheet(wb.create_sheet("Therapist Summary"), ["Therapist", *statuses, "Total"],
           [[name, *[c[s] for s in statuses], sum(c.values())] for name, c in sorted(summary.items())])
    if completion is not None:
        _sheet(wb.create_sheet("Slot Completion"),
               ["Date", "Therapist", "Therapy", "Working Slots", "Booked Slots", "Slots Done",
                "Pending", "Missed (therapist absent)", "Empty", "All Booked Slots Completed"],
               completion)
    return _respond(wb, filename)


def assignments_workbook(qs, filename):
    qs = qs.select_related("student__group", "therapy", "staff__user", "time_slot")
    wb = Workbook()
    ws = wb.active
    ws.title = "Slot Assignments"
    _sheet(ws, ["Student ID", "Student", "Group", "Therapy", "Therapist", "Time Slot",
                "Days", "From", "To", "Status"],
           [[a.student.student_id, a.student.name, a.student.group.name if a.student.group else "",
             a.therapy.name, a.staff.full_name, str(a.time_slot), a.weekday_display,
             a.start_date, a.end_date, "Current" if a.is_current else "Ended"] for a in qs])
    return _respond(wb, filename)


def attendance_workbook(student_qs, staff_qs, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Attendance"
    student_qs = student_qs.select_related("student").order_by("date", "student__name")
    _sheet(ws, ["Date", "Student ID", "Student", "Status", "Remarks"],
           [[a.date, a.student.student_id, a.student.name, a.get_status_display(), a.remarks]
            for a in student_qs])
    ws = wb.create_sheet("Staff Attendance")
    staff_qs = staff_qs.select_related("staff__user").order_by("date", "staff__employee_id")
    _sheet(ws, ["Date", "Employee ID", "Name", "Status", "Remarks"],
           [[a.date, a.staff.employee_id, a.staff.full_name, a.get_status_display(), a.remarks]
            for a in staff_qs])

    def summary(qs, ident, name):
        data = defaultdict(Counter)
        for a in qs:
            data[(ident(a), name(a))][a.get_status_display()] += 1
        return [[i, n, c["Present"], c["Absent"], c["On Leave"], sum(c.values())]
                for (i, n), c in sorted(data.items(), key=lambda x: x[0][1])]

    _sheet(wb.create_sheet("Student Summary"),
           ["Student ID", "Student", "Present", "Absent", "On Leave", "Total"],
           summary(student_qs, lambda a: a.student.student_id, lambda a: a.student.name))
    _sheet(wb.create_sheet("Staff Summary"),
           ["Employee ID", "Name", "Present", "Absent", "On Leave", "Total"],
           summary(staff_qs, lambda a: a.staff.employee_id, lambda a: a.staff.full_name))
    return _respond(wb, filename)


def person_workbook(kind, person, sessions, attendance, group_att, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    _sheet(ws, SESSION_HEADERS, _session_rows(
        sessions.select_related("student__group", "therapy", "staff__user", "time_slot")))
    _sheet(wb.create_sheet("Daily Attendance"), ["Date", "Status", "Remarks"],
           [[a.date, a.get_status_display(), a.remarks] for a in attendance])
    if kind == "student":
        _sheet(wb.create_sheet("Group Attendance"), ["Date", "Group", "Status", "Remarks"],
               [[g.group_session.date, g.group_session.group.name, g.get_status_display(), g.remarks]
                for g in group_att])
    return _respond(wb, filename)
